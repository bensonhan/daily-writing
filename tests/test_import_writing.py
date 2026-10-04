import tempfile
import unittest
from datetime import date
from pathlib import Path

from scripts.import_writing import (
    GENERATED_MARKER,
    WritingImportError,
    extract_text_upload,
    import_uploads,
    parse_date_heading,
    parse_entries,
    reflow_pdf_lines,
    render_markdown,
)


class ImportWritingTest(unittest.TestCase):
    def test_supported_date_headings(self):
        expected = date(2026, 10, 3)
        for heading in (
            "2026-10-03",
            "10/3/2026",
            "October 3, 2026",
            "Saturday, October 3, 2026",
            "October 3rd, 2026",
            "## October 3, 2026",
        ):
            with self.subTest(heading=heading):
                self.assertEqual(expected, parse_date_heading(heading))

    def test_entries_are_normalized_and_sorted_newest_first(self):
        paragraphs = [
            "October 2026",
            "October 1, 2026",
            "First paragraph.",
            "Second paragraph.",
            "10/3/2026",
            "Newest entry.",
        ]
        entries = parse_entries(paragraphs, 2026, 10)
        self.assertEqual(
            [date(2026, 10, 3), date(2026, 10, 1)],
            [entry.day for entry in entries],
        )
        self.assertEqual(
            f"{GENERATED_MARKER}\n\n# October 2026\n\n"
            "## 2026-10-03\n\nNewest entry.\n\n"
            "## 2026-10-01\n\nFirst paragraph.\n\nSecond paragraph.\n",
            render_markdown(entries, 2026, 10),
        )

    def test_pdf_visual_lines_are_reflowed(self):
        lines = [
            "October 2026",
            "",
            "October 3, 2026",
            "This sentence wrapped at the",
            "right edge of the PDF.",
            "",
            "This is a second paragraph.",
        ]
        self.assertEqual(
            [
                "October 2026",
                "October 3, 2026",
                "This sentence wrapped at the right edge of the PDF.",
                "This is a second paragraph.",
            ],
            reflow_pdf_lines(lines),
        )

    def test_rejects_wrong_month_and_unstructured_preamble(self):
        with self.assertRaisesRegex(WritingImportError, "does not match upload month"):
            parse_entries(["2026-09-30", "Wrong month."], 2026, 10)
        with self.assertRaisesRegex(WritingImportError, "text before the first date"):
            parse_entries(["Notes", "2026-10-01", "Entry."], 2026, 10)

    def test_rejects_empty_and_duplicate_entries(self):
        with self.assertRaisesRegex(WritingImportError, "entry has no writing"):
            parse_entries(["2026-10-01"], 2026, 10)
        with self.assertRaisesRegex(WritingImportError, "duplicate date"):
            parse_entries(
                ["2026-10-01", "One.", "October 1, 2026", "Two."],
                2026,
                10,
            )
        with self.assertRaisesRegex(WritingImportError, "invalid date heading"):
            parse_entries(["February 30, 2026", "Impossible."], 2026, 2)

    def test_imports_a_docx_end_to_end(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            uploads = root / "uploads"
            generated = root / ".generated-writing"
            writing = root / "writing"
            uploads.mkdir()
            writing.mkdir()

            document = Document()
            document.add_heading("October 2026", level=1)
            document.add_heading("October 3, 2026", level=2)
            document.add_paragraph("First paragraph.")
            document.add_paragraph("Second paragraph.")
            document.save(uploads / "2026-10.docx")

            written = import_uploads(uploads, generated, writing)

            self.assertEqual([generated / "2026-10.md"], written)
            self.assertEqual(
                f"{GENERATED_MARKER}\n\n# October 2026\n\n"
                "## 2026-10-03\n\nFirst paragraph.\n\nSecond paragraph.\n",
                written[0].read_text(encoding="utf-8"),
            )

    def test_imports_google_docs_markdown_and_preserves_formatting(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            uploads = root / "uploads"
            generated = root / ".generated-writing"
            writing = root / "writing"
            uploads.mkdir()
            writing.mkdir()
            source = uploads / "2026-10.md"
            source.write_text(
                "# October 2026\n\n"
                "## October 3, 2026\n\n"
                "A paragraph with **bold text**.\n\n"
                "- One\n- Two\n",
                encoding="utf-8",
            )

            written = import_uploads(uploads, generated, writing)

            self.assertEqual(
                f"{GENERATED_MARKER}\n\n# October 2026\n\n"
                "## 2026-10-03\n\nA paragraph with **bold text**.\n\n"
                "- One\n\n- Two\n",
                written[0].read_text(encoding="utf-8"),
            )

    def test_imports_google_docs_weekday_and_date_lines(self):
        paragraphs = [
            "Wednesday",
            "09/30/2026",
            "First paragraph.",
            "Second paragraph.",
            "Tuesday",
            "09/29/2026",
            "Earlier entry.",
        ]
        entries = parse_entries(paragraphs, 2026, 9)
        self.assertEqual(
            [date(2026, 9, 30), date(2026, 9, 29)],
            [entry.day for entry in entries],
        )
        self.assertEqual(
            ("First paragraph.", "Second paragraph."),
            entries[0].paragraphs,
        )

    def test_plain_text_uses_each_line_as_a_paragraph(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            source = Path(temporary_directory) / "2026-10.txt"
            source.write_text(
                "October 2026\nOctober 3, 2026\nFirst.\nSecond.\n",
                encoding="utf-8",
            )
            self.assertEqual(
                ["October 2026", "October 3, 2026", "First.", "Second."],
                extract_text_upload(source),
            )


if __name__ == "__main__":
    unittest.main()
